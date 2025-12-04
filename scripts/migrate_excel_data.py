#!/usr/bin/env python3
"""
ExcelデータをSQLite（D1ローカル）に移行するスクリプト
"""

import openpyxl
import sqlite3
import os
from datetime import datetime

# パス設定
EXCEL_PATH = '/home/user/uploaded_files/請求データ個人WEB部門026.xlsm'
DB_PATH = '/home/user/webapp/.wrangler/state/v3/d1/miniflare-D1DatabaseObject/local-dev-db.sqlite'

def get_db_connection():
    """データベース接続を取得"""
    # D1のローカルDBパスを探す
    base_path = '/home/user/webapp/.wrangler/state/v3/d1'
    if os.path.exists(base_path):
        for root, dirs, files in os.walk(base_path):
            for file in files:
                if file.endswith('.sqlite'):
                    return sqlite3.connect(os.path.join(root, file))
    raise Exception(f"SQLite database not found in {base_path}")

def migrate_company_info(wb, conn):
    """自社情報を移行"""
    print("=== 自社情報を移行中 ===")
    ws = wb['自社情報']
    
    cursor = conn.cursor()
    
    # 既存データを削除
    cursor.execute("DELETE FROM company_info")
    
    # データを読み込み
    data = {
        'company_name': ws['C3'].value or '',
        'department_name': ws['C4'].value or '',
        'person_name': ws['C5'].value or '',
        'postal_code': ws['C6'].value or '',
        'address': ws['C7'].value or '',
        'address_number': ws['C8'].value or '',
        'building_name': ws['C9'].value or '',
        'tel': ws['C10'].value or '',
        'fax': ws['C11'].value or '',
        'email': ws['C12'].value or '',
        'website': ws['C13'].value or '',
        'invoice_registration_no': ws['C14'].value or '',
        'bank_name': ws['C15'].value or '',
    }
    
    # 追加の銀行情報を読み込む（行が存在すれば）
    try:
        data['bank_branch'] = ws['C16'].value or ''
        data['account_type'] = ws['C17'].value or ''
        data['account_number'] = ws['C18'].value or ''
        data['account_holder'] = ws['C19'].value or ''
    except:
        data['bank_branch'] = ''
        data['account_type'] = ''
        data['account_number'] = ''
        data['account_holder'] = ''
    
    cursor.execute("""
        INSERT INTO company_info (
            company_name, department_name, person_name, postal_code,
            address, address_number, building_name, tel, fax, email,
            website, invoice_registration_no, bank_name, bank_branch,
            account_type, account_number, account_holder
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
    """, (
        data['company_name'], data['department_name'], data['person_name'],
        data['postal_code'], data['address'], data['address_number'],
        data['building_name'], data['tel'], data['fax'], data['email'],
        data['website'], data['invoice_registration_no'], data['bank_name'],
        data['bank_branch'], data['account_type'], data['account_number'],
        data['account_holder']
    ))
    
    conn.commit()
    print(f"  会社名: {data['company_name']}")
    print("  自社情報の移行完了")

def migrate_categories(wb, conn):
    """分類マスタを移行"""
    print("\n=== 分類マスタを移行中 ===")
    ws = wb['品目']
    
    cursor = conn.cursor()
    cursor.execute("DELETE FROM categories")
    
    # 3行目にカテゴリ名がある
    categories = set()
    for col in range(1, ws.max_column + 1):
        val = ws.cell(row=3, column=col).value
        if val and val.strip():
            categories.add(val.strip())
    
    count = 0
    for i, cat_name in enumerate(sorted(categories)):
        cursor.execute("""
            INSERT INTO categories (category_name, tax_rate, display_order)
            VALUES (?, ?, ?)
        """, (cat_name, 10, i + 1))
        count += 1
        print(f"  {cat_name}")
    
    conn.commit()
    print(f"  {count}件の分類を移行完了")

def migrate_products(wb, conn):
    """商品マスタを移行"""
    print("\n=== 商品マスタを移行中 ===")
    ws = wb['商品マスタ']
    
    cursor = conn.cursor()
    cursor.execute("DELETE FROM products")
    
    # カテゴリIDマッピングを取得
    cursor.execute("SELECT id, category_name FROM categories")
    category_map = {row[1]: row[0] for row in cursor.fetchall()}
    
    count = 0
    for row in range(2, ws.max_row + 1):
        product_name = ws.cell(row=row, column=1).value
        if not product_name:
            continue
        
        product_code = ws.cell(row=row, column=2).value or ''
        jan_code = ws.cell(row=row, column=3).value or ''
        unit_price = ws.cell(row=row, column=4).value or 0
        cost_price = ws.cell(row=row, column=5).value or 0
        category_name = ws.cell(row=row, column=7).value or ''
        min_lot = ws.cell(row=row, column=8).value or 1
        
        category_id = category_map.get(category_name.strip()) if category_name else None
        
        # unit_priceが文字列の場合は0にする
        try:
            unit_price = float(unit_price) if unit_price else 0
        except:
            unit_price = 0
        
        try:
            cost_price = float(cost_price) if cost_price else 0
        except:
            cost_price = 0
        
        cursor.execute("""
            INSERT INTO products (
                product_name, product_code, jan_code, category_id,
                unit_price, cost_price, retail_price, discount_rate, min_lot
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
        """, (
            product_name, product_code, str(jan_code), category_id,
            unit_price, cost_price, 0, 100, min_lot or 1
        ))
        count += 1
    
    conn.commit()
    print(f"  {count}件の商品を移行完了")

def migrate_clients(wb, conn):
    """取引先マスタを移行"""
    print("\n=== 取引先マスタを移行中 ===")
    ws = wb['取引先マスタ']
    
    cursor = conn.cursor()
    cursor.execute("DELETE FROM clients")
    
    count = 0
    for row in range(2, ws.max_row + 1):
        client_name = ws.cell(row=row, column=2).value
        if not client_name:
            continue
        
        # A列: №, B列: 取引先名, C列: 部署名, D列: 郵便番号, E列: 住所, 
        # F列: 住所番地, G列: ビル名等, H列: 取引先管理コード, I列: 担当者名,
        # J列: メールアドレス, K列: 締め日, L列: 支払い期日
        department = ws.cell(row=row, column=3).value or ''
        postal_code = ws.cell(row=row, column=4).value or ''
        address = ws.cell(row=row, column=5).value or ''
        address_number = ws.cell(row=row, column=6).value or ''
        building = ws.cell(row=row, column=7).value or ''
        client_code = ws.cell(row=row, column=8).value or ''
        person_name = ws.cell(row=row, column=9).value or ''
        email = ws.cell(row=row, column=10).value or ''
        closing_day = ws.cell(row=row, column=11).value or ''
        payment_day = ws.cell(row=row, column=12).value or ''
        
        cursor.execute("""
            INSERT INTO clients (
                client_name, department_name, postal_code, address,
                address_number, building_name, client_code, person_name,
                email, closing_day, payment_day
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """, (
            client_name, department, str(postal_code), address,
            str(address_number), building, str(client_code), person_name,
            email, str(closing_day), str(payment_day)
        ))
        count += 1
        print(f"  {client_name} (コード: {client_code})")
    
    conn.commit()
    print(f"  {count}件の取引先を移行完了")

def migrate_delivery_data(wb, conn):
    """納品データを移行"""
    print("\n=== 納品データを移行中 ===")
    ws = wb['納品データ']
    
    cursor = conn.cursor()
    cursor.execute("DELETE FROM delivery_items")
    cursor.execute("DELETE FROM deliveries")
    
    # 取引先とカテゴリのマッピング
    cursor.execute("SELECT id, client_name FROM clients")
    client_map = {row[1]: row[0] for row in cursor.fetchall()}
    
    cursor.execute("SELECT id, category_name FROM categories")
    category_map = {row[1]: row[0] for row in cursor.fetchall()}
    
    # 日付と取引先でグループ化して納品書を作成
    deliveries = {}
    
    for row in range(3, ws.max_row + 1):
        date_val = ws.cell(row=row, column=1).value
        client_name = ws.cell(row=row, column=2).value
        
        if not date_val or not client_name:
            continue
        
        # 日付を文字列に変換
        if isinstance(date_val, datetime):
            date_str = date_val.strftime('%Y-%m-%d')
        else:
            date_str = str(date_val)[:10]
        
        key = (date_str, client_name)
        if key not in deliveries:
            deliveries[key] = []
        
        product_name = ws.cell(row=row, column=3).value or ''
        quantity = ws.cell(row=row, column=4).value or 1
        unit_price = ws.cell(row=row, column=5).value or 0
        jan_code = ws.cell(row=row, column=6).value or ''
        category_name = ws.cell(row=row, column=7).value or ''
        
        try:
            quantity = int(quantity)
        except:
            quantity = 1
        
        try:
            unit_price = float(unit_price)
        except:
            unit_price = 0
        
        deliveries[key].append({
            'product_name': product_name,
            'quantity': quantity,
            'unit_price': unit_price,
            'jan_code': str(jan_code),
            'category_id': category_map.get(category_name.strip()) if category_name else None
        })
    
    # 納品書を作成
    delivery_count = 0
    item_count = 0
    
    for (date_str, client_name), items in deliveries.items():
        client_id = client_map.get(client_name)
        if not client_id:
            print(f"  警告: 取引先 '{client_name}' が見つかりません")
            continue
        
        # 取引先コードを取得
        cursor.execute("SELECT client_code FROM clients WHERE id = ?", (client_id,))
        client_code = cursor.fetchone()[0] or '000'
        
        # 納品書番号を生成
        year = date_str[2:4]
        month_day = date_str[5:7] + date_str[8:10]
        delivery_no = f"DS{year}-{client_code}-{month_day}"
        
        # 小計・税・合計を計算
        subtotal = sum(item['quantity'] * item['unit_price'] for item in items)
        tax_amount = int(subtotal * 0.1)
        total_amount = subtotal + tax_amount
        
        cursor.execute("""
            INSERT INTO deliveries (
                delivery_no, delivery_date, client_id, subtotal, tax_amount, total_amount, status
            ) VALUES (?, ?, ?, ?, ?, ?, ?)
        """, (delivery_no, date_str, client_id, subtotal, tax_amount, total_amount, 'issued'))
        
        delivery_id = cursor.lastrowid
        delivery_count += 1
        
        # 明細を追加
        for i, item in enumerate(items):
            cursor.execute("""
                INSERT INTO delivery_items (
                    delivery_id, product_name, jan_code, category_id,
                    quantity, unit_price, tax_rate, amount, display_order
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
            """, (
                delivery_id, item['product_name'], item['jan_code'], item['category_id'],
                item['quantity'], item['unit_price'], 10,
                item['quantity'] * item['unit_price'], i + 1
            ))
            item_count += 1
    
    conn.commit()
    print(f"  {delivery_count}件の納品書、{item_count}件の明細を移行完了")

def migrate_estimate_data(wb, conn):
    """見積データを移行"""
    print("\n=== 見積データを移行中 ===")
    ws = wb['見積書作成データ']
    
    cursor = conn.cursor()
    cursor.execute("DELETE FROM estimate_items")
    cursor.execute("DELETE FROM estimates")
    
    # 取引先マッピング
    cursor.execute("SELECT id, client_name FROM clients")
    client_map = {row[1]: row[0] for row in cursor.fetchall()}
    
    # 日付と取引先でグループ化
    estimates = {}
    
    for row in range(5, ws.max_row + 1):
        date_val = ws.cell(row=row, column=1).value
        client_name = ws.cell(row=row, column=2).value
        
        if not date_val or not client_name:
            continue
        
        if isinstance(date_val, datetime):
            date_str = date_val.strftime('%Y-%m-%d')
        else:
            date_str = str(date_val)[:10]
        
        key = (date_str, client_name)
        if key not in estimates:
            estimates[key] = []
        
        product_name = ws.cell(row=row, column=3).value or ''
        quantity = ws.cell(row=row, column=4).value or 1
        unit_price = ws.cell(row=row, column=5).value or 0
        jan_code = ws.cell(row=row, column=6).value or ''
        notes = ws.cell(row=row, column=7).value or ''
        
        try:
            quantity = int(quantity)
        except:
            quantity = 1
        
        try:
            unit_price = float(unit_price)
        except:
            unit_price = 0
        
        estimates[key].append({
            'product_name': product_name,
            'quantity': quantity,
            'unit_price': unit_price,
            'jan_code': str(jan_code),
            'notes': notes
        })
    
    # 見積書を作成
    estimate_count = 0
    item_count = 0
    
    for (date_str, client_name), items in estimates.items():
        client_id = client_map.get(client_name)
        if not client_id:
            print(f"  警告: 取引先 '{client_name}' が見つかりません")
            continue
        
        # 取引先コードを取得
        cursor.execute("SELECT client_code FROM clients WHERE id = ?", (client_id,))
        client_code = cursor.fetchone()[0] or '000'
        
        # 見積書番号を生成
        year = date_str[2:4]
        month_day = date_str[5:7] + date_str[8:10]
        estimate_no = f"ES{year}-{client_code}-{month_day}"
        
        # 小計・税・合計を計算
        subtotal = sum(item['quantity'] * item['unit_price'] for item in items)
        tax_amount = int(subtotal * 0.1)
        total_amount = subtotal + tax_amount
        
        cursor.execute("""
            INSERT INTO estimates (
                estimate_no, estimate_date, client_id, subtotal, tax_amount, total_amount, status
            ) VALUES (?, ?, ?, ?, ?, ?, ?)
        """, (estimate_no, date_str, client_id, subtotal, tax_amount, total_amount, 'draft'))
        
        estimate_id = cursor.lastrowid
        estimate_count += 1
        
        # 明細を追加
        for i, item in enumerate(items):
            cursor.execute("""
                INSERT INTO estimate_items (
                    estimate_id, product_name, jan_code,
                    quantity, unit_price, tax_rate, amount, notes, display_order
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
            """, (
                estimate_id, item['product_name'], item['jan_code'],
                item['quantity'], item['unit_price'], 10,
                item['quantity'] * item['unit_price'], item['notes'], i + 1
            ))
            item_count += 1
    
    conn.commit()
    print(f"  {estimate_count}件の見積書、{item_count}件の明細を移行完了")

def main():
    print("=" * 50)
    print("SmartBill データ移行スクリプト")
    print("=" * 50)
    
    # Excelファイルを読み込み
    print(f"\nExcelファイル読み込み中: {EXCEL_PATH}")
    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    print(f"シート一覧: {wb.sheetnames}")
    
    # データベース接続
    print("\nデータベース接続中...")
    conn = get_db_connection()
    print("接続成功")
    
    try:
        # 各データを移行
        migrate_company_info(wb, conn)
        migrate_categories(wb, conn)
        migrate_products(wb, conn)
        migrate_clients(wb, conn)
        migrate_delivery_data(wb, conn)
        migrate_estimate_data(wb, conn)
        
        print("\n" + "=" * 50)
        print("データ移行が完了しました！")
        print("=" * 50)
        
    finally:
        conn.close()
        wb.close()

if __name__ == '__main__':
    main()
